"""Resume text extraction shared by every endpoint that reads a resume.

Supports PDF (pypdf), DOCX (python-docx) and plain text. Scanned PDFs have no
text layer, so they are rejected with a clear message rather than returning
an empty string.
"""
from __future__ import annotations

import io

MAX_RESUME_BYTES = 5 * 1024 * 1024  # 5 MB
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")


class ResumeParseError(ValueError):
    """A problem the user can fix (wrong type, too big, no text)."""


def extract_text(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if not name.endswith(SUPPORTED_EXTENSIONS):
        raise ResumeParseError("Upload a PDF, DOCX or TXT file.")
    if len(content) > MAX_RESUME_BYTES:
        raise ResumeParseError("File is larger than 5 MB.")
    if not content:
        raise ResumeParseError("The file is empty.")

    if name.endswith(".txt"):
        text = content.decode("utf-8", errors="ignore")
    elif name.endswith(".docx"):
        text = _docx_text(content)
    else:
        text = _pdf_text(content)

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ResumeParseError(
            "No text found. If this is a scanned PDF, export it as a text PDF or DOCX."
        )
    return text


def _pdf_text(content: bytes) -> str:
    import pypdf

    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:  # pypdf raises many types for malformed files
        raise ResumeParseError(f"Could not read this PDF ({e.__class__.__name__}).") from e


def _docx_text(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover - dependency listed in requirements
        raise ResumeParseError("DOCX support is not installed on the server.") from e

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as e:
        raise ResumeParseError("Could not read this DOCX file.") from e
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
